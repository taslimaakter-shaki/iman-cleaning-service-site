import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


OUT_DIR = Path("/Users/taslimaaktershaki/Downloads/nyc_apartment_contacts_1000")
CONTACTS_PATH = OUT_DIR / "nyc_apartment_contacts.json"
XLSX_PATH = OUT_DIR / "nyc_apartment_contacts_1000.xlsx"
CSV_PATH = OUT_DIR / "nyc_apartment_contacts_1000.csv"


HEADERS = [
    "Name",
    "Category",
    "Address",
    "Phone",
    "Email",
    "All Emails Found",
    "Website",
    "Contact Page Checked",
    "Google Business Profile",
    "Source Query",
    "Scraped At",
]


FIELDS = [
    "name",
    "category",
    "address",
    "phone",
    "email",
    "all_emails_found",
    "website",
    "contact_page_checked",
    "google_business_profile",
    "source_query",
    "scraped_at",
]


def csv_escape(value):
    text = "" if value is None else str(value)
    return '"' + text.replace('"', '""') + '"'


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    contacts = json.loads(CONTACTS_PATH.read_text()) if CONTACTS_PATH.exists() else []

    wb = Workbook()
    ws = wb.active
    ws.title = "NYC Apartment Contacts"
    ws.append(HEADERS)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for item in contacts:
        ws.append([item.get(field, "") for field in FIELDS])

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    widths = {
        "A": 34,
        "B": 26,
        "C": 42,
        "D": 18,
        "E": 30,
        "F": 40,
        "G": 34,
        "H": 34,
        "I": 52,
        "J": 30,
        "K": 24,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    summary = wb.create_sheet("Summary")
    with_email = sum(1 for item in contacts if item.get("email"))
    with_phone = sum(1 for item in contacts if item.get("phone"))
    with_website = sum(1 for item in contacts if item.get("website"))
    summary_rows = [
        ("Total records", len(contacts)),
        ("Records with public email found", with_email),
        ("Records with phone", with_phone),
        ("Records with website", with_website),
        ("Output folder", str(OUT_DIR)),
    ]
    for row in summary_rows:
        summary.append(row)
    summary.column_dimensions["A"].width = 34
    summary.column_dimensions["B"].width = 80
    summary["A1"].font = Font(bold=True)

    wb.save(XLSX_PATH)

    lines = [",".join(csv_escape(header) for header in HEADERS)]
    for item in contacts:
        lines.append(",".join(csv_escape(item.get(field, "")) for field in FIELDS))
    CSV_PATH.write_text("\n".join(lines) + "\n")

    print(f"records={len(contacts)} emails={with_email} phones={with_phone} websites={with_website}")
    print(XLSX_PATH)
    print(CSV_PATH)


if __name__ == "__main__":
    main()
